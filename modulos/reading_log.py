"""Persistencia local de lecturas OCR y relaciones entre dos evidencias."""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import datetime
import math
from pathlib import Path

from config import ARCHIVO_REGISTRO_LECTURAS, ARCHIVO_REPORTE_LECTURAS


@dataclass(frozen=True)
class Reading:
    id: int
    created_at: str
    codigo: str
    texto_observado: str
    estado: str
    confianza: float
    score: float
    metodo: str
    foto_path: str
    reporte_path: str
    input_sha256: str
    relation_id: int | None
    relation_slot: int | None
    discarded_at: str | None


@dataclass(frozen=True)
class ReadingEntry:
    relation_id: int | None
    first: Reading
    second: Reading | None = None

    @property
    def desechado(self) -> bool:
        return self.first.discarded_at is not None


class ReadingLogStore:
    """Almacena lecturas sin borrar registros al reiniciar la aplicacion."""

    def __init__(self, db_path: str | Path = ARCHIVO_REGISTRO_LECTURAS):
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._initialize()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def _initialize(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS relaciones (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS lecturas (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL,
                    codigo TEXT NOT NULL DEFAULT '',
                    texto_observado TEXT NOT NULL DEFAULT '',
                    estado TEXT NOT NULL DEFAULT '',
                    confianza REAL NOT NULL DEFAULT 0,
                    score REAL NOT NULL DEFAULT 0,
                    metodo TEXT NOT NULL DEFAULT '',
                    foto_path TEXT NOT NULL DEFAULT '',
                    reporte_path TEXT NOT NULL DEFAULT '',
                    input_sha256 TEXT NOT NULL DEFAULT '',
                    relation_id INTEGER REFERENCES relaciones(id) ON DELETE SET NULL,
                    relation_slot INTEGER CHECK (relation_slot IN (1, 2)),
                    discarded_at TEXT
                );

                CREATE UNIQUE INDEX IF NOT EXISTS idx_lecturas_relation_slot
                ON lecturas(relation_id, relation_slot)
                WHERE relation_id IS NOT NULL;
                """
            )
            columns = {
                row[1]
                for row in connection.execute("PRAGMA table_info(lecturas)").fetchall()
            }
            if "discarded_at" not in columns:
                connection.execute("ALTER TABLE lecturas ADD COLUMN discarded_at TEXT")

    @staticmethod
    def _number(value) -> float:
        try:
            number = float(value)
            return number if math.isfinite(number) else 0.0
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _now() -> str:
        return datetime.now().astimezone().isoformat(timespec="seconds")

    def agregar_lectura(
        self,
        *,
        codigo: str = "",
        texto_observado: str = "",
        estado: str = "",
        confianza: float = 0.0,
        score: float = 0.0,
        metodo: str = "",
        foto_path: str = "",
        reporte_path: str = "",
        input_sha256: str = "",
        created_at: str | None = None,
    ) -> int:
        with self._connect() as connection:
            cursor = connection.execute(
                """
                INSERT INTO lecturas(
                    created_at, codigo, texto_observado, estado, confianza,
                    score, metodo, foto_path, reporte_path, input_sha256
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    created_at or self._now(),
                    str(codigo or "").strip(),
                    str(texto_observado or "").strip(),
                    str(estado or "").strip(),
                    self._number(confianza),
                    self._number(score),
                    str(metodo or "").strip(),
                    str(foto_path or ""),
                    str(reporte_path or ""),
                    str(input_sha256 or ""),
                ),
            )
            return int(cursor.lastrowid)

    def _row_to_reading(self, row: sqlite3.Row) -> Reading:
        return Reading(
            id=int(row["id"]),
            created_at=row["created_at"],
            codigo=row["codigo"],
            texto_observado=row["texto_observado"],
            estado=row["estado"],
            confianza=float(row["confianza"]),
            score=float(row["score"]),
            metodo=row["metodo"],
            foto_path=row["foto_path"],
            reporte_path=row["reporte_path"],
            input_sha256=row["input_sha256"],
            relation_id=(int(row["relation_id"]) if row["relation_id"] is not None else None),
            relation_slot=(int(row["relation_slot"]) if row["relation_slot"] is not None else None),
            discarded_at=row["discarded_at"],
        )

    def _readings(self, include_discarded: bool = False) -> list[Reading]:
        with self._connect() as connection:
            query = "SELECT * FROM lecturas"
            if not include_discarded:
                query += " WHERE discarded_at IS NULL"
            query += " ORDER BY id DESC"
            rows = connection.execute(query).fetchall()
        return [self._row_to_reading(row) for row in rows]

    def listar_filas(self, include_discarded: bool = False) -> list[ReadingEntry]:
        """Devuelve filas visuales: una lectura sola o dos relacionadas."""
        singles: list[ReadingEntry] = []
        grouped: dict[int, dict[int, Reading]] = {}
        for reading in self._readings(include_discarded=include_discarded):
            if reading.relation_id is None:
                singles.append(ReadingEntry(None, reading))
            else:
                grouped.setdefault(reading.relation_id, {})[
                    reading.relation_slot or 1
                ] = reading

        entries = list(singles)
        for relation_id, slots in grouped.items():
            first = slots.get(1) or slots.get(2)
            second = slots.get(2) if 1 in slots else None
            if first is not None:
                entries.append(ReadingEntry(relation_id, first, second))

        entries.sort(
            key=lambda entry: max(
                entry.first.id,
                entry.second.id if entry.second is not None else entry.first.id,
            ),
            reverse=True,
        )
        return entries

    def relacionar(self, first_id: int, second_id: int) -> int:
        if first_id == second_id:
            raise ValueError("Selecciona dos lecturas distintas")
        with self._connect() as connection:
            ids = {int(first_id), int(second_id)}
            placeholders = ",".join("?" for _ in ids)
            rows = connection.execute(
                f"SELECT id, relation_id FROM lecturas WHERE id IN ({placeholders})",
                tuple(ids),
            ).fetchall()
            if len(rows) != 2:
                raise ValueError("Una de las lecturas ya no existe")
            if any(row["relation_id"] is not None for row in rows):
                raise ValueError("Solo se pueden relacionar lecturas sin relacion")

            relation = connection.execute(
                "INSERT INTO relaciones(created_at) VALUES (?)",
                (self._now(),),
            )
            relation_id = int(relation.lastrowid)
            connection.execute(
                "UPDATE lecturas SET relation_id=?, relation_slot=1 WHERE id=?",
                (relation_id, int(first_id)),
            )
            connection.execute(
                "UPDATE lecturas SET relation_id=?, relation_slot=2 WHERE id=?",
                (relation_id, int(second_id)),
            )
            return relation_id

    def separar(self, relation_id: int) -> None:
        with self._connect() as connection:
            connection.execute(
                "UPDATE lecturas SET relation_id=NULL, relation_slot=NULL WHERE relation_id=?",
                (int(relation_id),),
            )
            connection.execute("DELETE FROM relaciones WHERE id=?", (int(relation_id),))

    def cambiar_desechado(
        self,
        *,
        reading_id: int,
        relation_id: int | None,
        discarded: bool,
    ) -> None:
        """Marca una lectura o una relacion completa sin borrar sus datos."""
        timestamp = self._now() if discarded else None
        with self._connect() as connection:
            if relation_id is not None:
                connection.execute(
                    "UPDATE lecturas SET discarded_at=? WHERE relation_id=?",
                    (timestamp, int(relation_id)),
                )
            else:
                connection.execute(
                    "UPDATE lecturas SET discarded_at=? WHERE id=?",
                    (timestamp, int(reading_id)),
                )

    def total_lecturas(self) -> int:
        with self._connect() as connection:
            return int(connection.execute("SELECT COUNT(*) FROM lecturas").fetchone()[0])

    def exportar_excel(self, destination: str | Path = ARCHIVO_REPORTE_LECTURAS) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise RuntimeError("Falta instalar openpyxl para exportar el registro") from exc

        path = Path(destination)
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Lecturas"
        headers = [
            "Relacion",
            "Estado registro",
            "ID lectura 1",
            "Codigo 1",
            "Texto observado 1",
            "Fecha 1",
            "Confianza 1",
            "Score 1",
            "Metodo 1",
            "Foto 1",
            "Reporte OCR 1",
            "ID lectura 2",
            "Codigo 2",
            "Texto observado 2",
            "Fecha 2",
            "Confianza 2",
            "Score 2",
            "Metodo 2",
            "Foto 2",
            "Reporte OCR 2",
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="45475A")
            cell.alignment = Alignment(horizontal="center")

        def path_cell(value: str):
            return Path(value).name if value else ""

        def add_link(cell, value: str) -> None:
            if not value:
                return
            try:
                cell.hyperlink = Path(value).resolve().as_uri()
                cell.style = "Hyperlink"
            except (OSError, ValueError):
                return

        def values(reading: Reading | None) -> list:
            if reading is None:
                return ["", "", "", "", "", "", "", "", ""]
            return [
                reading.id,
                reading.codigo,
                reading.texto_observado,
                reading.created_at,
                reading.confianza,
                reading.score,
                reading.metodo,
                path_cell(reading.foto_path),
                path_cell(reading.reporte_path),
            ]

        entries = self.listar_filas(include_discarded=True)
        for entry in entries:
            row = [
                entry.relation_id or "",
                "Desechado" if entry.desechado else "Activo",
            ] + values(entry.first) + values(entry.second)
            sheet.append(row)
            excel_row = sheet.max_row
            add_link(sheet.cell(excel_row, 10), entry.first.foto_path)
            add_link(sheet.cell(excel_row, 11), entry.first.reporte_path)
            if entry.second is not None:
                add_link(sheet.cell(excel_row, 19), entry.second.foto_path)
                add_link(sheet.cell(excel_row, 20), entry.second.reporte_path)
            for column in (7, 8, 16, 17):
                sheet.cell(excel_row, column).number_format = "0.00"
            for column in (7, 16):
                sheet.cell(excel_row, column).number_format = "0.0%"

        widths = [12, 16, 12, 24, 32, 24, 14, 12, 24, 30, 30, 12, 24, 32, 24, 14, 12, 24, 30, 30]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index) if index <= 26 else f"A{index}"].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        summary = workbook.create_sheet("Resumen")
        summary.append(["Indicador", "Valor"])
        summary.append(["Filas exportadas", len(entries)])
        summary.append(["Lecturas guardadas", self.total_lecturas()])
        summary.append(["Filas desechadas", sum(entry.desechado for entry in entries)])
        summary.append(["Generado", self._now()])
        summary.column_dimensions["A"].width = 24
        summary.column_dimensions["B"].width = 32

        workbook.save(path)
        return path
