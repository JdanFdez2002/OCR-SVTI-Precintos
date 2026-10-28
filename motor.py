"""Herramienta independiente para evaluar el motor OCR de la aplicacion.

La ventana reutiliza ``OcrWorker`` y el lector visual compartido de
``modulos.ocr_reader``. No contiene una segunda implementacion de OCR ni
aplica preprocesamiento adicional a las imagenes seleccionadas.
"""

from __future__ import annotations

import sys
import json
import time
from pathlib import Path

import cv2
from PySide6.QtCore import Qt, QTimer
from PySide6.QtGui import QImage, QPixmap
from PySide6.QtWidgets import (
    QApplication,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMainWindow,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from modulos.ocr_worker import OcrWorker
from ui.styles import APP_STYLESHEET


EXTENSIONES_IMAGEN = {
    ".bmp",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}


class MotorWindow(QMainWindow):
    """Ventana para procesar una carpeta de imagenes con el OCR existente."""

    COLUMNAS = [
        "Imagen",
        "Texto detectado",
        "Confianza",
        "Score",
        "Metodo",
        "Tiempo",
        "Estado",
    ]

    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Evaluacion del motor OCR - SVTI")
        self.resize(1280, 760)
        self.setStyleSheet(APP_STYLESHEET)

        self._imagenes: list[Path] = []
        self._indice_actual = -1
        self._inicio_imagen = 0.0
        self._procesando = False
        self._ocr_listo = False

        self._ocr = OcrWorker(self)
        self._ocr.engine_ready.connect(self._on_ocr_listo)
        self._ocr.engine_failed.connect(self._on_ocr_fallo_motor)
        self._ocr.result.connect(self._on_ocr_resultado)
        self._ocr.analyzed.connect(self._on_analisis_detallado)
        self._ocr.failed.connect(self._on_ocr_error)
        self._ocr.start()
        self._ocr.pedir_warmup()

        self._armar_ui()

    def _armar_ui(self) -> None:
        self._btn_carpeta = QPushButton("Seleccionar carpeta")
        self._btn_carpeta.clicked.connect(self._seleccionar_carpeta)

        self._btn_procesar = QPushButton("Procesar carpeta")
        self._btn_procesar.setObjectName("primario")
        self._btn_procesar.clicked.connect(self._procesar_carpeta)
        self._btn_procesar.setEnabled(False)

        self._btn_exportar = QPushButton("Exportar Excel")
        self._btn_exportar.clicked.connect(self._exportar_excel)
        self._btn_exportar.setEnabled(False)

        self._lbl_motor = QLabel("Motor OCR: cargando...")
        self._lbl_carpeta = QLabel("Carpeta: ninguna")
        self._lbl_carpeta.setWordWrap(True)
        self._lbl_estado = QLabel("Selecciona una carpeta con imagenes.")
        self._lbl_progreso = QLabel("0 imagenes")

        self._barra = QProgressBar()
        self._barra.setRange(0, 1)
        self._barra.setValue(0)

        controles = QHBoxLayout()
        controles.addWidget(self._btn_carpeta)
        controles.addWidget(self._btn_procesar)
        controles.addWidget(self._btn_exportar)
        controles.addStretch(1)

        encabezado = QVBoxLayout()
        encabezado.addLayout(controles)
        encabezado.addWidget(self._lbl_motor)
        encabezado.addWidget(self._lbl_carpeta)
        encabezado.addWidget(self._lbl_estado)
        encabezado.addWidget(self._barra)
        encabezado.addWidget(self._lbl_progreso)

        self._tabla = QTableWidget(0, len(self.COLUMNAS))
        self._tabla.setHorizontalHeaderLabels(self.COLUMNAS)
        self._tabla.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self._tabla.setSelectionMode(
            QTableWidget.SelectionMode.SingleSelection
        )
        self._tabla.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self._tabla.verticalHeader().setVisible(False)
        self._tabla.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents
        )
        self._tabla.horizontalHeader().setStretchLastSection(True)
        self._tabla.itemSelectionChanged.connect(self._mostrar_fila_seleccionada)

        self._vista = QLabel("Selecciona una fila para ver la imagen.")
        self._vista.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._vista.setMinimumSize(360, 300)
        self._vista.setStyleSheet(
            "border: 1px solid #45475A; background: #11111B;"
        )
        self._lbl_detalle = QLabel("")
        self._lbl_detalle.setWordWrap(True)

        caja_vista = QGroupBox("Imagen seleccionada")
        lay_vista = QVBoxLayout(caja_vista)
        lay_vista.addWidget(self._vista, 1)
        lay_vista.addWidget(self._lbl_detalle)

        divisor = QSplitter(Qt.Orientation.Horizontal)
        divisor.addWidget(self._tabla)
        divisor.addWidget(caja_vista)
        divisor.setStretchFactor(0, 3)
        divisor.setStretchFactor(1, 2)

        central = QWidget()
        lay = QVBoxLayout(central)
        lay.addLayout(encabezado)
        lay.addWidget(divisor, 1)
        self.setCentralWidget(central)

    def _seleccionar_carpeta(self) -> None:
        if self._procesando:
            return
        carpeta = QFileDialog.getExistingDirectory(
            self,
            "Seleccionar carpeta de imagenes",
        )
        if carpeta:
            self._cargar_carpeta(Path(carpeta))

    def _cargar_carpeta(self, carpeta: Path) -> None:
        try:
            imagenes = sorted(
                (
                    ruta
                    for ruta in carpeta.rglob("*")
                    if ruta.is_file()
                    and ruta.suffix.lower() in EXTENSIONES_IMAGEN
                ),
                key=lambda ruta: str(ruta).lower(),
            )
        except OSError as exc:
            QMessageBox.warning(
                self,
                "Carpeta",
                f"No se pudo leer la carpeta seleccionada:\n{exc}",
            )
            return

        self._imagenes = imagenes
        self._indice_actual = -1
        self._tabla.clearContents()
        self._tabla.setRowCount(len(imagenes))
        self._btn_exportar.setEnabled(False)
        self._lbl_carpeta.setText(f"Carpeta: {carpeta}")
        self._barra.setRange(0, len(imagenes) or 1)
        self._barra.setValue(0)
        self._actualizar_progreso(0)

        for fila, ruta in enumerate(imagenes):
            self._tabla.setItem(fila, 0, QTableWidgetItem(ruta.name))
            self._tabla.item(fila, 0).setToolTip(str(ruta))
            self._tabla.setItem(fila, 1, QTableWidgetItem(""))
            self._tabla.setItem(fila, 2, QTableWidgetItem(""))
            self._tabla.setItem(fila, 3, QTableWidgetItem(""))
            self._tabla.setItem(fila, 4, QTableWidgetItem(""))
            self._tabla.setItem(fila, 5, QTableWidgetItem(""))
            self._tabla.setItem(fila, 6, QTableWidgetItem("Pendiente"))

        self._vista.setText("Selecciona una fila para ver la imagen.")
        self._vista.setPixmap(QPixmap())
        self._lbl_detalle.clear()
        self._btn_procesar.setEnabled(bool(imagenes) and self._ocr_listo)
        if imagenes:
            if self._ocr_listo:
                self._lbl_estado.setText(
                    f"{len(imagenes)} imagenes encontradas. Iniciando procesamiento..."
                )
                self._procesar_carpeta()
            else:
                self._lbl_estado.setText(
                    f"{len(imagenes)} imagenes encontradas. Esperando al motor OCR."
                )
        else:
            self._lbl_estado.setText(
                "No se encontraron imagenes compatibles en la carpeta."
            )

    def _procesar_carpeta(self) -> None:
        if self._procesando or not self._imagenes:
            return
        if not self._ocr_listo:
            QMessageBox.warning(self, "OCR", "El motor OCR aun no esta listo.")
            return

        self._procesando = True
        self._indice_actual = 0
        self._btn_carpeta.setEnabled(False)
        self._btn_procesar.setEnabled(False)
        self._barra.setRange(0, len(self._imagenes))
        self._barra.setValue(0)
        self._actualizar_progreso(0)
        self._enviar_siguiente()

    def _enviar_siguiente(self) -> None:
        while self._indice_actual < len(self._imagenes):
            fila = self._indice_actual
            ruta = self._imagenes[fila]
            self._tabla.selectRow(fila)
            self._establecer_estado(fila, "Procesando...")
            self._lbl_estado.setText(f"Procesando: {ruta.name}")

            try:
                imagen = cv2.imread(str(ruta), cv2.IMREAD_COLOR)
            except Exception as exc:
                self._registrar_error(fila, f"No se pudo leer: {exc}")
                self._indice_actual += 1
                self._actualizar_progreso(self._indice_actual)
                continue

            if imagen is None or imagen.size == 0:
                self._registrar_error(fila, "No se pudo leer la imagen")
                self._indice_actual += 1
                self._actualizar_progreso(self._indice_actual)
                continue

            self._inicio_imagen = time.perf_counter()
            # Se entrega el ndarray BGR directamente al mismo worker que usa
            # la aplicacion de camara; el motor decide su propio pipeline.
            self._ocr.pedir_ocr(imagen)
            return

        self._finalizar_procesamiento()

    def _on_ocr_resultado(
        self, texto: str, score: float, variante: str, conf: float
    ) -> None:
        if not self._procesando or not self._indice_valido():
            return

        fila = self._indice_actual
        tiempo = time.perf_counter() - self._inicio_imagen
        self._tabla.item(fila, 1).setText(texto)
        self._tabla.item(fila, 1).setToolTip(texto)
        if texto:
            self._tabla.item(fila, 2).setText(f"{conf * 100:.1f} %")
            self._tabla.item(fila, 3).setText(f"{score:.2f}")
        else:
            self._tabla.item(fila, 2).setText("—")
            self._tabla.item(fila, 3).setText("—")
        self._tabla.item(fila, 4).setText(variante)
        self._tabla.item(fila, 5).setText(f"{tiempo:.3f} s")
        detail = self._tabla.item(fila,0).data(Qt.ItemDataRole.UserRole) or {}
        self._establecer_estado(fila, detail.get('status','texto_reconocido' if texto else 'sin_texto'))

        self._indice_actual += 1
        self._actualizar_progreso(self._indice_actual)
        self._enviar_siguiente()

    def _on_ocr_error(self, mensaje: str) -> None:
        if not self._procesando or not self._indice_valido():
            return

        fila = self._indice_actual
        tiempo = time.perf_counter() - self._inicio_imagen
        self._tabla.item(fila, 5).setText(f"{tiempo:.3f} s")
        self._registrar_error(fila, f"Error OCR: {mensaje}")
        self._indice_actual += 1
        self._actualizar_progreso(self._indice_actual)
        self._enviar_siguiente()

    def _on_ocr_listo(self) -> None:
        self._ocr_listo = True
        self._lbl_motor.setText("Motor OCR: listo")
        if self._imagenes and not self._procesando:
            self._procesar_carpeta()
        else:
            self._btn_procesar.setEnabled(bool(self._imagenes))

    def _on_ocr_fallo_motor(self, mensaje: str) -> None:
        self._ocr_listo = False
        self._lbl_motor.setText(f"Motor OCR: error - {mensaje}")
        self._btn_procesar.setEnabled(False)
        QMessageBox.critical(self, "OCR", f"No se pudo cargar el motor OCR:\n{mensaje}")

    def _exportar_excel(self) -> None:
        filas = self._filas_exportables()
        if not filas:
            QMessageBox.information(
                self,
                "Exportar resultados",
                "Todavia no hay resultados procesados para exportar.",
            )
            return

        ruta_sugerida = self._imagenes[filas[0]].parent / "resultados_ocr.xlsx"
        ruta, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar resultados OCR",
            str(ruta_sugerida),
            "Archivos Excel (*.xlsx)",
        )
        if not ruta:
            return

        destino = Path(ruta)
        if destino.suffix.lower() != ".xlsx":
            destino = destino.with_suffix(".xlsx")

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill

            libro = Workbook()
            hoja = libro.active
            hoja.title = "Resultados OCR"

            encabezados = [
                "Nombre de archivo",
                "Ruta de la imagen",
                "Texto reconocido",
                "Confianza",
                "Score",
                "Metodo",
                "Tiempo de procesamiento",
                "Estado",
            ]
            hoja.append(encabezados)

            for celda in hoja[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="45475A")
                celda.alignment = Alignment(horizontal="center")

            for fila in filas:
                ruta_imagen = self._imagenes[fila]
                hoja.append(
                    [
                        ruta_imagen.name,
                        str(ruta_imagen),
                        self._valor_exportable(fila, 1),
                        self._valor_exportable(fila, 2),
                        self._valor_exportable(fila, 3),
                        self._valor_exportable(fila, 4),
                        self._valor_exportable(fila, 5),
                        self._valor_celda(fila, 6),
                    ]
                )

            anchos = [
                28,
                75,
                32,
                14,
                14,
                30,
                24,
                35,
            ]
            for indice, ancho in enumerate(anchos, start=1):
                hoja.column_dimensions[chr(64 + indice)].width = ancho
            for fila_excel in hoja.iter_rows(min_row=2):
                for celda in fila_excel:
                    celda.alignment = Alignment(vertical="top", wrap_text=True)

            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
            evidencia = destino.with_name(destino.stem + f'_evidence_{time.time_ns()}.json')
            with evidencia.open('x',encoding='utf-8') as stream:
                json.dump({'kind':'visual_ocr_evidence_no_manifest',
                           'records':[{'image':str(self._imagenes[fila]),
                                       'analysis':self._tabla.item(fila,0).data(Qt.ItemDataRole.UserRole)}
                                      for fila in filas]},stream,ensure_ascii=False,indent=2,allow_nan=False)
            libro.save(str(destino))
        except ImportError:
            QMessageBox.critical(
                self,
                "Exportar resultados",
                "No se encontro openpyxl en el entorno virtual. No se creo ningun archivo.",
            )
            return
        except Exception as exc:
            QMessageBox.critical(
                self,
                "Exportar resultados",
                f"No se pudo guardar el archivo Excel:\n{exc}",
            )
            return

        QMessageBox.information(
            self,
            "Exportar resultados",
            f"Se exportaron {len(filas)} resultados a:\n{destino}\nEvidencia OCR:\n{evidencia}",
        )

    def _filas_exportables(self) -> list[int]:
        filas = []
        for fila in range(self._tabla.rowCount()):
            estado = self._valor_celda(fila, 6)
            if estado and estado not in {"Pendiente", "Procesando..."}:
                filas.append(fila)
        return filas

    def _valor_celda(self, fila: int, columna: int) -> str:
        item = self._tabla.item(fila, columna)
        return item.text().strip() if item is not None else ""

    def _valor_exportable(self, fila: int, columna: int) -> str:
        valor = self._valor_celda(fila, columna)
        if valor in {"", "—", "-"}:
            return ""
        return valor

    def _mostrar_fila_seleccionada(self) -> None:
        filas = self._tabla.selectionModel().selectedRows()
        if not filas:
            return
        fila = filas[0].row()
        if fila < 0 or fila >= len(self._imagenes):
            return

        ruta = self._imagenes[fila]
        try:
            imagen = cv2.imread(str(ruta), cv2.IMREAD_COLOR)
        except Exception as exc:
            self._vista.setText(f"No se pudo cargar la imagen:\n{exc}")
            return
        if imagen is None or imagen.size == 0:
            self._vista.setText("No se pudo cargar la imagen.")
            return

        alto, ancho = imagen.shape[:2]
        rgb = cv2.cvtColor(imagen, cv2.COLOR_BGR2RGB)
        qimage = QImage(
            rgb.data,
            ancho,
            alto,
            rgb.strides[0],
            QImage.Format.Format_RGB888,
        ).copy()
        pixmap = QPixmap.fromImage(qimage)
        self._vista.setPixmap(
            pixmap.scaled(
                self._vista.size(),
                Qt.AspectRatioMode.KeepAspectRatio,
                Qt.TransformationMode.SmoothTransformation,
            )
        )
        self._lbl_detalle.setText(f"{ruta.name}\n{ancho} x {alto} px")

    def _registrar_error(self, fila: int, mensaje: str) -> None:
        self._tabla.item(fila, 4).setText("—")
        self._tabla.item(fila, 5).setText(
            self._tabla.item(fila, 5).text() or "—"
        )
        self._establecer_estado(fila, mensaje)

    def _establecer_estado(self, fila: int, estado: str) -> None:
        self._tabla.item(fila, 6).setText(estado)
        self._tabla.item(fila, 6).setToolTip(estado)
        self._btn_exportar.setEnabled(bool(self._filas_exportables()))

    def _actualizar_progreso(self, completadas: int) -> None:
        total = len(self._imagenes)
        self._barra.setValue(completadas)
        self._lbl_progreso.setText(f"{completadas} de {total} imagenes procesadas")

    def _finalizar_procesamiento(self) -> None:
        self._procesando = False
        self._btn_carpeta.setEnabled(True)
        self._btn_procesar.setEnabled(bool(self._imagenes) and self._ocr_listo)
        self._lbl_estado.setText("Procesamiento finalizado. Puedes repetirlo.")

    def _indice_valido(self) -> bool:
        return 0 <= self._indice_actual < len(self._imagenes)

    def _on_analisis_detallado(self, packet):
        analysis = packet['analysis']
        if self._indice_valido():
            item = self._tabla.item(self._indice_actual, 0)
            if item is not None:
                item.setData(Qt.ItemDataRole.UserRole, analysis.to_dict())
                item.setToolTip(analysis.status + '\n' + analysis.reason + '\nCandidatos:\n' +
                                '\n'.join(c.text for c in analysis.candidates))

    def closeEvent(self, event) -> None:
        self.centralWidget().setEnabled(False)
        self._ocr.detener()
        if self._ocr.isRunning():
            event.ignore()
            QTimer.singleShot(100, self.close)
        else:
            event.accept()


def main() -> int:
    app = QApplication(sys.argv)
    app.setOrganizationName("ACCELERA")
    app.setApplicationName("SVTI-Motor-OCR")
    ventana = MotorWindow()
    ventana.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
