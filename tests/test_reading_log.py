import tempfile
import unittest
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from modulos.reading_log import ReadingLogStore


class ReadingLogStoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db_path = Path(self.tmp.name) / "registro.sqlite3"

    def test_readings_survive_a_new_store_instance(self):
        first = ReadingLogStore(self.db_path)
        reading_id = first.agregar_lectura(
            codigo="ASDASD",
            texto_observado="ASDASD",
            confianza=.91,
            foto_path=str(Path(self.tmp.name) / "foto_1.png"),
        )

        reopened = ReadingLogStore(self.db_path)
        entries = reopened.listar_filas()

        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0].first.id, reading_id)
        self.assertEqual(entries[0].first.codigo, "ASDASD")
        self.assertIsNone(entries[0].relation_id)

    def test_related_readings_are_exported_in_one_row_with_photo_links(self):
        store = ReadingLogStore(self.db_path)
        photo_1 = Path(self.tmp.name) / "foto_1.png"
        photo_2 = Path(self.tmp.name) / "foto_2.png"
        photo_1.write_bytes(b"fixture")
        photo_2.write_bytes(b"fixture")
        first_id = store.agregar_lectura(codigo="ASDASD", foto_path=str(photo_1))
        second_id = store.agregar_lectura(codigo="ASJDJAD", foto_path=str(photo_2))

        relation_id = store.relacionar(first_id, second_id)
        entries = store.listar_filas()
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0].relation_id, relation_id)
        self.assertEqual(entries[0].first.codigo, "ASDASD")
        self.assertEqual(entries[0].second.codigo, "ASJDJAD")

        destination = Path(self.tmp.name) / "reporte.xlsx"
        store.exportar_excel(destination)
        workbook = load_workbook(destination)
        sheet = workbook["Lecturas"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 4).value, "ASDASD")
        self.assertEqual(sheet.cell(2, 13).value, "ASJDJAD")
        self.assertIsNotNone(sheet.cell(2, 10).hyperlink)
        self.assertIsNotNone(sheet.cell(2, 19).hyperlink)

    def test_separating_relation_returns_two_independent_rows(self):
        store = ReadingLogStore(self.db_path)
        first_id = store.agregar_lectura(codigo="A")
        second_id = store.agregar_lectura(codigo="B")
        relation_id = store.relacionar(first_id, second_id)

        store.separar(relation_id)
        entries = store.listar_filas()

        self.assertEqual(len(entries), 2)
        self.assertTrue(all(entry.relation_id is None for entry in entries))

    def test_discarding_a_relation_moves_both_readings_without_deleting_them(self):
        store = ReadingLogStore(self.db_path)
        first_id = store.agregar_lectura(codigo="ASDASD")
        second_id = store.agregar_lectura(codigo="ASJDJAD")
        relation_id = store.relacionar(first_id, second_id)

        store.cambiar_desechado(
            reading_id=first_id,
            relation_id=relation_id,
            discarded=True,
        )

        self.assertEqual(store.listar_filas(), [])
        discarded = store.listar_filas(include_discarded=True)
        self.assertEqual(len(discarded), 1)
        self.assertTrue(discarded[0].desechado)
        self.assertEqual(discarded[0].first.codigo, "ASDASD")
        self.assertEqual(discarded[0].second.codigo, "ASJDJAD")

        store.cambiar_desechado(
            reading_id=first_id,
            relation_id=relation_id,
            discarded=False,
        )
        recovered = store.listar_filas()
        self.assertEqual(len(recovered), 1)
        self.assertFalse(recovered[0].desechado)

    def test_existing_database_is_migrated_without_losing_readings(self):
        with sqlite3.connect(self.db_path) as connection:
            connection.executescript(
                """
                CREATE TABLE relaciones (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE lecturas (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL,
                    codigo TEXT NOT NULL DEFAULT '',
                    texto_observado TEXT NOT NULL DEFAULT '',
                    estado TEXT NOT NULL DEFAULT '',
                    confianza REAL NOT NULL DEFAULT 0,
                    score REAL NOT NULL DEFAULT 0,
                    metodo TEXT NOT NULL DEFAULT '',
                    foto_path TEXT NOT NULL DEFAULT '',
                    reporte_path TEXT NOT NULL DEFAULT '',
                    input_sha256 TEXT NOT NULL DEFAULT '',
                    relation_id INTEGER,
                    relation_slot INTEGER
                );
                INSERT INTO lecturas(created_at, codigo) VALUES ('2026-01-01T10:00:00-03:00', 'LEGACY');
                """
            )

        store = ReadingLogStore(self.db_path)
        entries = store.listar_filas()
        self.assertEqual(entries[0].first.codigo, "LEGACY")
        store.cambiar_desechado(
            reading_id=entries[0].first.id,
            relation_id=None,
            discarded=True,
        )
        self.assertEqual(store.listar_filas(), [])
        self.assertEqual(len(store.listar_filas(include_discarded=True)), 1)


if __name__ == "__main__":
    unittest.main()
